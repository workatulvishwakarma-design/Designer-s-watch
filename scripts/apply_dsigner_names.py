import openpyxl
import json
import re

# 1. Parse Excel sheet
wb = openpyxl.load_workbook(r'public/images/new-img/Collection Names Selected (2).xlsx')
sheet = wb["D'signer"]

model_to_collection = {}

for r in range(3, sheet.max_row + 1):
    cname = sheet.cell(row=r, column=1).value
    if not cname: continue
    cname = str(cname).strip()
    for c in range(3, sheet.max_column + 1):
        v = sheet.cell(row=r, column=c).value
        if v is not None:
            v_str = str(v).strip()
            # If compound string like '820G(gents) & 850L (ladies)'
            if '&' in v_str:
                for num in re.findall(r'\d+', v_str):
                    model_to_collection[num] = cname
                for code in re.findall(r'(\d+[A-Za-z]*)', v_str):
                    model_to_collection[code.upper()] = cname
            else:
                for num in re.findall(r'\d+', v_str):
                    model_to_collection[num] = cname
                clean_v = re.sub(r'^(Grandeur|Tactix|Stratos)\s*', '', v_str, flags=re.I).strip()
                if clean_v:
                    model_to_collection[clean_v.upper()] = cname
                model_to_collection[v_str.upper()] = cname

# Codebase known mappings
model_to_collection['915'] = 'Grandeur'
model_to_collection['875'] = 'Tactix'
model_to_collection['980'] = 'Grandeur'
model_to_collection['J905'] = 'Echo'

def get_collection(model_no, series):
    m_clean = str(model_no).strip().upper()
    s_clean = str(series).strip().upper()
    
    # 1. Exact model match
    if m_clean in model_to_collection:
        return model_to_collection[m_clean]
    # 2. Exact series match
    if s_clean in model_to_collection:
        return model_to_collection[s_clean]
    # 3. Numeric series match
    s_num = re.search(r'\d+', s_clean)
    if s_num and s_num.group(0) in model_to_collection:
        return model_to_collection[s_num.group(0)]
    # 4. Model prefix digits
    m_num = re.search(r'\d+', m_clean)
    if m_num and m_num.group(0) in model_to_collection:
        return model_to_collection[m_num.group(0)]
    
    return "Classic"

# 2. Update dsigner_master_inventory.json
inventory_path = 'src/data/dsigner_master_inventory.json'
with open(inventory_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

men_items = data.get('men', [])
women_items = data.get('women', [])

print(f"Updating {len(men_items)} Men items and {len(women_items)} Women items...")

for item in men_items:
    col = get_collection(item['modelNo'], item['series'])
    item['collection'] = col
    item['name'] = f"{col} {item['modelNo']}"

for item in women_items:
    col = get_collection(item['modelNo'], item['series'])
    item['collection'] = col
    item['name'] = f"{col} {item['modelNo']}"

# Write back
with open(inventory_path, 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, ensure_ascii=False)

print("dsigner_master_inventory.json successfully updated.")

# Verify sample
print("\nSample verification (Men):")
for item in men_items[:5]:
    print(f"  modelNo: {item['modelNo']} -> name: {item['name']}, collection: {item['collection']}")

print("\nSample verification (Women):")
for item in women_items[:5]:
    print(f"  modelNo: {item['modelNo']} -> name: {item['name']}, collection: {item['collection']}")
